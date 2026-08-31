import os
import pytest
from fpdf import FPDF
import docx
from app.parsers.pdf_parser import extract_text_from_pdf
from app.parsers.docx_parser import extract_text_from_docx
from app.parsers.resume_parser import clean_text, detect_sections, parse_resume
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

@pytest.fixture
def dummy_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="JOHN DOE", ln=True)
    pdf.cell(200, 10, txt="EXPERIENCE", ln=True)
    pdf.cell(200, 10, txt="Software Engineer", ln=True)
    
    file_path = "test_resume.pdf"
    pdf.output(file_path)
    yield file_path
    if os.path.exists(file_path):
        os.remove(file_path)

@pytest.fixture
def dummy_docx():
    doc = docx.Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Computer Science Degree")
    
    file_path = "test_resume.docx"
    doc.save(file_path)
    yield file_path
    if os.path.exists(file_path):
        os.remove(file_path)

def test_extract_text_from_pdf(dummy_pdf):
    text = extract_text_from_pdf(dummy_pdf)
    assert "JOHN DOE" in text
    assert "EXPERIENCE" in text

def test_extract_text_from_docx(dummy_docx):
    text = extract_text_from_docx(dummy_docx)
    assert "JANE DOE" in text
    assert "EDUCATION" in text

def test_clean_text():
    raw_text = "   This \r\n is \n\n\n a   test.   "
    cleaned = clean_text(raw_text)
    assert "   " not in cleaned
    assert "\n\n\n" not in cleaned
    assert cleaned.startswith("This")
    assert cleaned.endswith("test.")

def test_detect_sections():
    text = "JOHN DOE\n\nEXPERIENCE\nWorked as a developer.\n\nSKILLS\nPython, React"
    sections = detect_sections(text)
    assert "Worked as a developer." in sections["experience"]
    assert "Python, React" in sections["skills"]
    assert "JOHN DOE" in sections["other"]

def test_parse_resume_full(dummy_pdf):
    result = parse_resume(dummy_pdf, "PDF")
    assert result["character_count"] > 0
    assert result["word_count"] > 0
    assert "JOHN DOE" in result["full_text"]
    assert "Software Engineer" in result["sections"]["experience"]

def test_extract_text_from_docx_with_tables():
    doc = docx.Document()
    doc.add_paragraph("ALICE DEVELOPER")
    # Add a table (common in resumes)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].text = "TECHNICAL SKILLS"
    table.cell(0, 1).paragraphs[0].text = "Python, FastAPI, Docker, SQL"
    table.cell(1, 0).paragraphs[0].text = "EXPERIENCE"
    table.cell(1, 1).paragraphs[0].text = "Built web apps with React and TypeScript."
    
    file_path = "test_table_resume.docx"
    doc.save(file_path)
    try:
        text = extract_text_from_docx(file_path)
        assert "ALICE DEVELOPER" in text
        assert "Python" in text
        assert "FastAPI" in text
        assert "Docker" in text
        assert "React" in text
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)

def test_parse_endpoint(dummy_docx):
    with open(dummy_docx, "rb") as f:
        response = client.post(
            "/api/resume/parse",
            files={"file": ("test_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["file_type"] == "DOCX"
    assert data["character_count"] > 0
    assert "JANE DOE" in data["full_text"]
